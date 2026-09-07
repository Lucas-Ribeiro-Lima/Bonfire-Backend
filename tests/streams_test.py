from io import BytesIO
from unittest.mock import MagicMock

import pandas as pd
import pytest
from docx import Document

from domain.entities import AutoInfracao
from infrastructure.parsers.exceptions import (
    DocumentReadError,
    IncorrectInstanceError,
    InvalidDocumentDataError,
    QuantityOfAtasMismatchError,
)
from infrastructure.parsers.pyingestion.streams import (
    BonfireInfracaoWriteStream,
    BonfireRecursoWriteStream,
    InfracoesCsvInputStream,
    InfracoesTransformStream,
    InfracoesXlsInputStream,
    RecursosDocxInputStream,
    SanitizedTextIO,
)

# ==========================================
# WRITE STREAMS TESTS
# ==========================================


def test_bonfire_recurso_write_stream_primeira_instancia():
    mock_db_manager = MagicMock()
    mock_session = mock_db_manager.session.return_value.__enter__.return_value
    mock_repo = mock_session.get_recurso_repository.return_value
    mock_repo.insert_primeira_instancia.return_value = 2

    stream = BonfireRecursoWriteStream(
        mock_db_manager, first_instance=True, batch_size=2
    )

    # Write a single item and a list of items
    stream.write({"NUM_RECURSO": "1"})
    stream.write([{"NUM_RECURSO": "2"}])
    stream.flush()

    assert mock_repo.insert_primeira_instancia.call_count >= 1


def test_bonfire_recurso_write_stream_segunda_instancia():
    mock_db_manager = MagicMock()
    mock_session = mock_db_manager.session.return_value.__enter__.return_value
    mock_repo = mock_session.get_recurso_repository.return_value
    mock_repo.insert_segunda_instancia.return_value = 1

    stream = BonfireRecursoWriteStream(
        mock_db_manager, first_instance=False, batch_size=10
    )
    stream.write({"NUM_RECURSO": "10"})
    stream.flush()

    mock_repo.insert_segunda_instancia.assert_called_once()


def test_bonfire_infracao_write_stream():
    mock_db_manager = MagicMock()
    mock_session = mock_db_manager.session.return_value.__enter__.return_value
    mock_repo = mock_session.get_autoinfracao_repository.return_value
    mock_repo.insert_bulk.return_value = 1

    stream = BonfireInfracaoWriteStream(mock_db_manager, batch_size=5)
    stream.write(
        {
            "NUM_AI": "12345-A",
            "NUM_NOTF": "NOTF-12345",
            "TIP_PENL": "MULTA",
            "NOM_CONC": "Consórcio BH Leste",
            "COD_LINH": "61",
            "NOM_LINH": "Estação Vilarinho",
            "DAT_OCOR_INFR": "2026-08-28T09:30:00",
            "COD_IRRG_FISC": 101,
            "ARTIGO": "Art. 1",
            "QTE_PONT": 3,
            "DAT_EMIS_NOTF": "2026-08-28T10:00:00",
            "DAT_LIMT_RECU": "2026-09-28T10:00:00",
            "VAL_INFR": 150.50,
        }
    )
    stream.flush()

    mock_repo.insert_bulk.assert_called_once()
    called_arg = mock_repo.insert_bulk.call_args[0][0]
    assert len(called_arg) == 1
    assert isinstance(called_arg[0], AutoInfracao)
    assert called_arg[0].notice_number == "12345-A"


# ==========================================
# SANITIZED TEXT IO TESTS
# ==========================================


def test_sanitized_text_io_accents_and_empty():
    empty_stream = BytesIO(b"")
    sanitized_empty = SanitizedTextIO(empty_stream)
    assert sanitized_empty.read() == ""

    # Stream with Portuguese accents and latin-1 chars
    raw_data = "Ação de Trânsito & Concessão".encode("utf-8")
    binary_stream = BytesIO(raw_data)
    sanitized = SanitizedTextIO(binary_stream)

    result = sanitized.read()
    assert result == "Acao de Transito & Concessao"

    sanitized.seek(0)
    assert binary_stream.tell() == 0


# ==========================================
# CSV & XLS INPUT STREAM TESTS
# ==========================================


def test_infracoes_csv_input_stream_detect_separator():
    stream = InfracoesCsvInputStream()

    pipe_data = BytesIO(b"COL1|COL2|COL3\nVAL1|VAL2|VAL3")
    assert stream._detect_separator(pipe_data) == "|"

    comma_data = BytesIO(b"COL1,COL2,COL3\nVAL1,VAL2,VAL3")
    assert stream._detect_separator(comma_data) == ","

    # Broken stream fallback
    broken_stream = MagicMock()
    broken_stream.read.side_effect = Exception("Read error")
    assert stream._detect_separator(broken_stream) == ";"


def test_infracoes_csv_input_stream_read():
    csv_bytes = b"NUM_AI;DAT_OCOR_INFR\n12345-A;01/01/2026\n"
    stream = InfracoesCsvInputStream()

    chunks = list(stream.read(BytesIO(csv_bytes)))
    assert len(chunks) == 1
    assert "NUM_AI" in chunks[0].columns

    # Test error wrapping
    bad_stream = MagicMock()
    bad_stream.read.side_effect = RuntimeError("Fatal disk error")
    with pytest.raises(DocumentReadError):
        list(stream.read(bad_stream))


def test_infracoes_xls_input_stream_error():
    stream = InfracoesXlsInputStream()
    with pytest.raises(DocumentReadError):
        list(stream.read(BytesIO(b"not an excel file")))


# ==========================================
# TRANSFORM STREAM TESTS
# ==========================================


def test_infracoes_transform_missing_dat_limt_recu():
    df = pd.DataFrame(
        {
            "NUM_AI": ["12345-A"],
            "DAT_LIMT_RECU": [""],
        }
    )
    stream = InfracoesTransformStream(
        datetime_format="%d/%m/%Y", date_format="%d/%m/%Y"
    )

    with pytest.raises(InvalidDocumentDataError) as exc_info:
        stream.transform(df)
    assert "DAT_LIMT_RECU" in str(exc_info.value)


# ==========================================
# DOCX INPUT STREAM TESTS
# ==========================================


def test_docx_stream_segunda_instancia_with_atas_raises_incorrect_instance():
    doc = Document()
    doc.add_paragraph(
        "PUBLICADO NO DIARIO OFICIAL DO MUNICIPIO DE BELO HORIZONTE EM 01/01/2026"
    )
    doc.add_paragraph("ATA DA 10ª SESSÃO")

    stream = RecursosDocxInputStream(first_instance=False)
    with pytest.raises(IncorrectInstanceError):
        stream.extract_atas(doc)


def test_docx_stream_atas_tables_mismatch():
    doc = Document()
    doc.add_paragraph("ATA DA 10ª SESSÃO")
    doc.add_paragraph("ATA DA 11ª SESSÃO")
    # 2 atas found, but 0 tables exist in doc
    stream = RecursosDocxInputStream(first_instance=True)
    with pytest.raises(QuantityOfAtasMismatchError):
        stream.extract_atas(doc)


def test_docx_stream_empty_num_recurso():
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    table.cell(0, 0).text = "RECURSO"
    table.cell(0, 1).text = "AUTO DE INFRAÇÃO"
    table.cell(0, 2).text = "RECORRENTE"
    table.cell(0, 3).text = "DECISÃO"

    # Row with empty recurso number
    table.cell(1, 0).text = ""
    table.cell(1, 1).text = "12345A"
    table.cell(1, 2).text = "Consorcio"
    table.cell(1, 3).text = "PROCEDENTE"

    stream = RecursosDocxInputStream(first_instance=True)
    with pytest.raises(InvalidDocumentDataError) as exc_info:
        list(stream.process_table(table, "2026-01-01", 1))
    assert "número do recurso está vazio" in str(exc_info.value)
